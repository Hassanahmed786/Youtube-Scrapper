"""Flask app: paste YouTube URLs, scrape them, then save to Google Sheets or download."""
import csv
import io
import logging
import os
import secrets
import uuid
from collections import OrderedDict
from datetime import datetime

from flask import Flask, Response, jsonify, render_template, request, session
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font

import sheets_client
import youtube_client
from utils import COLUMNS, check_config, load_config, parse_urls, result_to_row, setup_logging

# read .env before anything else uses the settings
load_config()
setup_logging()
log = logging.getLogger("app")

MAX_URLS = 50  # per request
MAX_STORED_SESSIONS = 200  # keep memory small
XLSX_CELL_LIMIT = 32767  # Excel cell size limit
NO_RESULTS_MESSAGE = "There are no results yet. Paste some URLs and click Scrape first."

app = Flask(__name__)
# if no secret key is set, make a random one (sessions reset on restart)
app.config["SECRET_KEY"] = os.getenv("FLASK_SECRET_KEY") or secrets.token_hex(32)
app.config["MAX_CONTENT_LENGTH"] = 1024 * 1024  # 1 MB is plenty for 50 URLs
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"

# The cookie is too small for all the results (descriptions can be long),
# so the cookie only holds an ID and the results live here in memory.
LATEST_RESULTS = OrderedDict()


# ---------- helpers ----------

def error_response(message, status=400):
    """JSON error that the page shows to the user."""
    return jsonify({"error": message}), status


def store_results(results):
    """Remember the latest results of this browser session."""
    if "sid" not in session:
        session["sid"] = uuid.uuid4().hex
    sid = session["sid"]
    LATEST_RESULTS[sid] = results
    LATEST_RESULTS.move_to_end(sid)
    while len(LATEST_RESULTS) > MAX_STORED_SESSIONS:
        LATEST_RESULTS.popitem(last=False)  # drop the oldest session


def load_results():
    """Get the latest results of this browser session (empty list if none)."""
    return LATEST_RESULTS.get(session.get("sid"), [])


def file_name(extension):
    """A file name like youtube_videos_20240501_123000.csv"""
    return f"youtube_videos_{datetime.now().strftime('%Y%m%d_%H%M%S')}.{extension}"


def download_response(data, mimetype, extension):
    """Wrap file bytes in a response that the browser downloads."""
    return Response(
        data,
        mimetype=mimetype,
        headers={"Content-Disposition": f'attachment; filename="{file_name(extension)}"'},
    )


def csv_safe(value):
    """Stop Excel from running text that starts with = + - @ as a formula."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@"):
        return "'" + value
    return value


def xlsx_text(value):
    """Remove characters Excel cannot store and cut to the Excel cell limit."""
    return ILLEGAL_CHARACTERS_RE.sub("", value)[:XLSX_CELL_LIMIT]


# ---------- routes ----------

@app.get("/")
def index():
    return render_template("index.html", max_urls=MAX_URLS)


@app.post("/scrape")
def scrape():
    data = request.get_json(silent=True) or {}
    text = data.get("urls", request.form.get("urls", ""))
    if isinstance(text, list):
        text = "\n".join(str(item) for item in text)
    if not isinstance(text, str):
        text = ""

    urls = parse_urls(text)  # no blank lines, no duplicates
    if not urls:
        return error_response("Please paste at least one YouTube URL or video ID.")
    if len(urls) > MAX_URLS:
        return error_response(f"Too many URLs ({len(urls)}). The limit is {MAX_URLS} per request.")

    try:
        results, errors = youtube_client.get_videos(urls)
    except Exception:
        log.exception("Scrape crashed")
        return error_response("Something went wrong while scraping. See the server console for details.", 500)

    store_results(results)
    return jsonify({"results": results, "errors": errors})


@app.post("/save-to-sheet")
def save_to_sheet():
    results = load_results()
    if not results:
        return error_response(NO_RESULTS_MESSAGE)

    try:
        summary = sheets_client.save_results(results)
    except sheets_client.SheetsError as err:
        log.error("Save failed: %s", err)
        return error_response(str(err))
    except Exception:
        log.exception("Save crashed")
        return error_response("Something went wrong while saving. See the server console for details.", 500)

    summary["sheet_url"] = sheets_client.sheet_url()
    return jsonify(summary)


@app.get("/download/csv")
def download_csv():
    results = load_results()
    if not results:
        return error_response(NO_RESULTS_MESSAGE)

    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(COLUMNS)
    for result in results:
        writer.writerow([csv_safe(value) for value in result_to_row(result)])

    # "utf-8-sig" adds the BOM so Excel shows non-English titles correctly
    data = buffer.getvalue().encode("utf-8-sig")
    return download_response(data, "text/csv", "csv")


@app.get("/download/xlsx")
def download_xlsx():
    results = load_results()
    if not results:
        return error_response(NO_RESULTS_MESSAGE)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Videos"

    for col, name in enumerate(COLUMNS, start=1):
        sheet.cell(row=1, column=col, value=name).font = Font(bold=True)

    for row_number, result in enumerate(results, start=2):
        for col, value in enumerate(result_to_row(result), start=1):
            cell = sheet.cell(row=row_number, column=col)
            if isinstance(value, str):
                cell.value = xlsx_text(value)
                cell.data_type = "s"  # always text, never a formula
            else:
                cell.value = value

    # make the file nicer to open
    sheet.freeze_panes = "A2"
    for letter, width in zip("ABCDEFG", (45, 25, 60, 22, 12, 45, 22)):
        sheet.column_dimensions[letter].width = width

    buffer = io.BytesIO()
    workbook.save(buffer)
    return download_response(
        buffer.getvalue(),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "xlsx",
    )


if __name__ == "__main__":
    # debug mode starts the app twice, only check the config in the first process
    if os.environ.get("WERKZEUG_RUN_MAIN") != "true":
        check_config()
    app.run(debug=True, port=5000)
